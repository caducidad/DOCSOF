#!/usr/bin/env python3
"""
build_entregable.py  —  genera un .docx a partir de la plantilla SOFIA/OPTIMA
y los ficheros de borrador de un entregable.

Uso:
  python build_entregable.py <carpeta_entregable> <plantilla.docx> <salida.docx>
"""
import sys, os, glob, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def build(entregable_dir, plantilla_docx, output_docx):
    entregable_dir = os.path.abspath(entregable_dir)
    doc = Document(plantilla_docx)
    body = doc.element.body
    for el in [c for c in body if c.tag != qn('w:sectPr')]:
        body.remove(el)

    def _runs(para, text):
        for m in re.finditer(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|[^*`]+)', text):
            chunk = m.group(0)
            if not chunk: continue
            if chunk.startswith('**') and chunk.endswith('**'):
                para.add_run(chunk[2:-2]).bold = True
            elif chunk.startswith('*') and chunk.endswith('*'):
                para.add_run(chunk[1:-1]).italic = True
            elif chunk.startswith('`') and chunk.endswith('`'):
                r = para.add_run(chunk[1:-1])
                r.font.name = 'Courier New'; r.font.size = Pt(9)
            elif '[PENDIENTE' in chunk:
                r = para.add_run(chunk)
                r.font.color.rgb = RGBColor(0xCC, 0, 0); r.bold = True
            else:
                if chunk: para.add_run(chunk)

    def add_title(text):
        # Título principal del documento → estilo Title de la plantilla
        p = doc.add_paragraph(text, style='Title'); return p

    def add_subtitle(text):
        # Subtítulo (cabecera PT/línea/tarea)
        p = doc.add_paragraph(style='Subtitle')
        _runs(p, text); return p

    def add_heading(text, level):
        # H1 en MD → Heading 1 (secciones numeradas: 1. Introducción)
        # H2 en MD → Heading 2 (subsecciones: 1.1. Objeto)
        # H3 en MD → Heading 3
        # H4 en MD → Heading 4
        style = f'Heading {min(max(level, 1), 4)}'
        return doc.add_paragraph(text, style=style)

    def add_normal(text):
        p = doc.add_paragraph(style='Normal'); _runs(p, text); return p

    def add_caption(text):
        p = doc.add_paragraph(style='Caption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text).italic = True; return p

    def add_image(rel_path, caption):
        abs_path = os.path.normpath(os.path.join(entregable_dir, rel_path))
        if not os.path.exists(abs_path):
            add_normal(f'[IMAGEN NO ENCONTRADA: {rel_path}]'); return
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(abs_path, width=Inches(5.5))
        except Exception as e:
            add_normal(f'[ERROR IMAGEN: {e}]'); return
        if caption:
            add_caption(caption)

    def add_table_md(rows_lines):
        rows = [l for l in rows_lines if not re.match(r'^\s*\|[-: |]+\|\s*$', l)]
        if not rows: return
        cols = [c.strip() for c in rows[0].strip().strip('|').split('|')]
        ncols = len(cols)
        tbl = doc.add_table(rows=len(rows), cols=ncols)
        try: tbl.style = 'Table Grid'
        except: pass
        for ri, row in enumerate(rows):
            cells = [c.strip() for c in row.strip().strip('|').split('|')]
            for ci in range(ncols):
                cell_txt = cells[ci] if ci < len(cells) else ''
                tc = tbl.rows[ri].cells[ci]
                tc.text = ''
                p = tc.paragraphs[0]
                if ri == 0: p.add_run(cell_txt).bold = True
                else: _runs(p, cell_txt)

    # ── parse MD ───────────────────────────────────────────────────────────────
    files = sorted(glob.glob(os.path.join(entregable_dir, 'borrador/*.md')))
    combined = '\n\n'.join(open(f, encoding='utf-8').read() for f in files)
    combined = re.sub(r'\(\.\./(material/capturas/)', r'(\1', combined)

    lines = combined.splitlines()
    i = 0; table_buf = []; in_code = False; code_buf = []
    # Track if we're in the first file (cabecera) for Title/Subtitle handling
    cabecera_done = False

    def flush_table():
        if table_buf: add_table_md(table_buf); table_buf.clear()

    while i < len(lines):
        line = lines[i]

        # fenced code
        if line.strip().startswith('```'):
            flush_table(); in_code = not in_code
            if not in_code and code_buf:
                p = doc.add_paragraph(style='Normal')
                p.add_run('\n'.join(code_buf)).font.name = 'Courier New'
                code_buf.clear()
            i += 1; continue
        if in_code: code_buf.append(line); i += 1; continue

        # table
        if line.strip().startswith('|'): table_buf.append(line); i += 1; continue
        else: flush_table()

        # heading
        m = re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            level = len(m.group(1)); text = m.group(2)
            if level == 1 and not cabecera_done:
                # First H1 = document title
                add_title(text); cabecera_done = True
            else:
                add_heading(text, level)
            i += 1; continue

        # bold paragraph immediately after title = subtitle (PT3 - L3.5...)
        if not cabecera_done:
            i += 1; continue  # skip anything before title

        # subtitle: **text** on its own line
        m2 = re.match(r'^\s*\*\*([^*]+)\*\*\s*$', line)
        if m2 and cabecera_done and i < 20:  # only near top for subtitle detection
            add_subtitle(m2.group(1)); i += 1; continue

        # image  ![caption](path)
        m = re.match(r'^\s*!\[([^\]]*)\]\(([^\)]+)\)', line)
        if m:
            path = m.group(2)
            alt = m.group(1)
            # Look ahead: next non-blank line starting with *Figura → use as caption, skip it
            j = i + 1
            while j < len(lines) and not lines[j].strip(): j += 1
            if j < len(lines) and re.match(r'^\s*\*Figura', lines[j]):
                caption = lines[j].strip().strip('*')
                i = j  # skip the caption line
            else:
                caption = alt
            add_image(path, caption); i += 1; continue

        # standalone italic lines (captions not after images)
        m = re.match(r'^\s*\*([^*].+[^*])\*\s*$', line)
        if m: add_caption(m.group(1)); i += 1; continue

        # bullet
        m = re.match(r'^\s*[-*]\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Paragraph')
            _runs(p, m.group(1)); i += 1; continue

        # numbered list
        m = re.match(r'^\s*\d+\.\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Paragraph')
            _runs(p, m.group(1)); i += 1; continue

        # blank
        if not line.strip(): i += 1; continue

        # normal
        add_normal(line); i += 1

    flush_table()
    os.makedirs(os.path.dirname(os.path.abspath(output_docx)), exist_ok=True)
    doc.save(output_docx)
    print(f'OK: {output_docx}  ({len(doc.paragraphs)} párrafos, {len(files)} secciones)')

if __name__ == '__main__':
    if len(sys.argv) != 4:
        print(__doc__); sys.exit(1)
    build(sys.argv[1], sys.argv[2], sys.argv[3])
